"""Excel file semantics; dataset ownership stays in the host."""
from contextlib import ExitStack
from datetime import date, datetime, time, timedelta
import hashlib
import json
import math
import re
from taskweave.plugins.sdk import PluginError
from .files import check, path_for, file_bytes, load, publish, write_guard


def worksheet(book, name=None, create=False):
    if name is None:
        return book.active
    if not isinstance(name, str) or not name or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name):
        raise PluginError('EXCEL_SHEET_INVALID')
    if name in book.sheetnames:
        return book[name]
    if create:
        if name.lower() in {n.lower() for n in book.sheetnames}:
            raise PluginError('EXCEL_SHEET_INVALID', '工作表名称大小写冲突')
        return book.create_sheet(name)
    raise PluginError('EXCEL_SHEET_MISSING', '工作表不存在')


def inspect(ctx, inputs):
    check(ctx)
    root, path = path_for(ctx, inputs['path'], existing=True)
    book = load(file_bytes(path))
    try:
        return {'path': path.relative_to(root).as_posix(),
                'sheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in book.worksheets]}
    finally:
        book.close()


def read(ctx, inputs):
    from openpyxl.utils.cell import range_boundaries
    check(ctx)
    root, path = path_for(ctx, inputs['path'], existing=True)
    content = file_bytes(path)
    with ExitStack() as stack:
        book = load(content)
        stack.callback(book.close)
        sheet = worksheet(book, inputs.get('sheet'))
        cached = None
        if inputs.get('formula_mode', 'values') == 'values':
            cache_book = load(content, values=True)
            stack.callback(cache_book.close)
            cached = cache_book[sheet.title]
        try:
            bounds = range_boundaries(inputs['range']) if inputs.get('range') else (1, 1, sheet.max_column or 1, sheet.max_row or 1)
            left, top, right, bottom = bounds
            if any(type(v) is not int or v < 1 for v in bounds) or left > right or top > bottom:
                raise ValueError()
            width, height = right - left + 1, bottom - top + 1
            if width > 256 or height > 100001 or width * height > 1000256:
                raise PluginError('EXCEL_LIMIT', '读取区域超过数据集上限，请指定更小区域')
        except (ValueError, TypeError) as exc:
            raise PluginError('EXCEL_RANGE_INVALID') from exc
        cell_rows = sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right)
        cached_rows = cached.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right) if cached else None
        rows, size = [], 0
        for cells in cell_rows:
            check(ctx)
            cache = next(cached_rows) if cached_rows else None
            values = []
            for index, cell in enumerate(cells):
                if cell.data_type == 'e':
                    raise PluginError('EXCEL_CELL_ERROR', '读取区域包含错误单元格')
                value = cell.value
                if cell.data_type == 'f' and cached_rows:
                    value = cache[index].value
                    if value is None:
                        raise PluginError('EXCEL_FORMULA_CACHE_MISSING', '公式没有缓存值；请先在表格软件中计算保存，或读取公式文本')
                    if cache[index].data_type == 'e':
                        raise PluginError('EXCEL_CELL_ERROR')
                if isinstance(value, (datetime, date, time)):
                    value = value.isoformat()
                elif isinstance(value, timedelta):
                    value = str(value)
                values.append(value)
            try:
                size += len(json.dumps(values, ensure_ascii=False, allow_nan=False).encode('utf-8'))
            except (TypeError, ValueError) as exc:
                raise PluginError('EXCEL_VALUE_INVALID') from exc
            if size > 16 * 1024 * 1024:
                raise PluginError('EXCEL_LIMIT', '读取数据超过 16 MiB')
            rows.append(values)
        columns = inputs.get('columns')
        header = inputs.get('header', True)
        if header:
            headings = rows.pop(0)
            if columns is None:
                if any(not isinstance(v, str) or not v.strip() for v in headings):
                    raise PluginError('EXCEL_HEADER_INVALID', '表头必须为非空文本，或提供 columns')
                columns = [{'name': v, 'type': 'any'} for v in headings]
        elif columns is None:
            raise PluginError('EXCEL_HEADER_INVALID', 'header=false 时需提供 columns')
        if len(columns) != width:
            raise PluginError('EXCEL_HEADER_INVALID', '列声明与区域宽度不一致')
        ref = ctx.datasets.create(inputs.get('name', sheet.title), columns, rows)
        return {'path': path.relative_to(root).as_posix(), 'sheet': sheet.title, 'dataset': ref, 'row_count': len(rows)}


def matrix(ctx, inputs):
    if 'values' in inputs and 'dataset' in inputs:
        raise PluginError('EXCEL_INPUT_INVALID', 'values 与 dataset 只选其一')
    if 'dataset' in inputs:
        description = ctx.datasets.describe(inputs['dataset'])
        table = ctx.datasets.read(inputs['dataset'], 0, 100000)
        if table['truncated']:
            raise PluginError('EXCEL_LIMIT', '拒绝导出被截断的数据集')
        names = [c['name'] for c in description['columns']]
        rows = [[row[key] for key in names] for row in table['rows']]
        return ([names] if inputs.get('include_header', True) else []) + rows
    return inputs.get('values', [])


def write_cells(ctx, sheet, values, start_cell, formulas):
    from openpyxl.utils.cell import coordinate_to_tuple
    if not isinstance(start_cell, str) or not re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]*', start_cell):
        raise PluginError('EXCEL_RANGE_INVALID')
    row_start, col_start = coordinate_to_tuple(start_cell)
    if not isinstance(values, list) or len(values) > 100001:
        raise PluginError('EXCEL_LIMIT')
    width = len(values[0]) if values and isinstance(values[0], (list, tuple)) else 0
    if (width > 256 or width * len(values) > 1000256 or row_start + max(len(values), 1) - 1 > 1048576
            or col_start + max(width, 1) - 1 > 16384):
        raise PluginError('EXCEL_LIMIT')
    size = 0
    for row_index, row in enumerate(values, row_start):
        check(ctx)
        if not isinstance(row, (list, tuple)) or len(row) != width:
            raise PluginError('EXCEL_INPUT_INVALID', '写入数据必须是等宽二维数组')
        try:
            size += len(json.dumps(row, ensure_ascii=False, allow_nan=False).encode('utf-8'))
        except (TypeError, ValueError) as exc:
            raise PluginError('EXCEL_VALUE_INVALID') from exc
        if size > 16 * 1024 * 1024:
            raise PluginError('EXCEL_LIMIT', '写入数据超过 16 MiB')
        for column_index, value in enumerate(row, col_start):
            if value is not None and type(value) not in (str, int, float, bool):
                raise PluginError('EXCEL_VALUE_INVALID')
            if ((type(value) is float and not math.isfinite(value))
                    or (type(value) is int and abs(value) > 999999999999999)
                    or (isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)))):
                raise PluginError('EXCEL_VALUE_INVALID', '数值精度、文本长度或字符无效；长编号请用文本')
            cell = sheet.cell(row_index, column_index)
            try:
                cell.value = value
                if isinstance(value, str) and not (formulas and value.startswith('=')):
                    cell.data_type = 's'
            except (ValueError, AttributeError) as exc:
                raise PluginError('EXCEL_CELL_WRITE_FAILED', '目标单元格无法写入，可能位于合并区域') from exc
    return len(values), width


def create(ctx, inputs):
    from openpyxl import Workbook
    check(ctx)
    root, path = path_for(ctx, inputs['path'], writing=True)
    if path.exists():
        raise PluginError('EXCEL_TARGET_EXISTS')
    book = Workbook()
    try:
        sheet = book.active
        if inputs.get('sheet'):
            new_sheet = worksheet(book, inputs['sheet'], create=True)
            if new_sheet is not sheet:
                book.remove(sheet)
                sheet = new_sheet
        rows, columns = write_cells(ctx, sheet, matrix(ctx, inputs), inputs.get('start_cell', 'A1'), inputs.get('formulas', False))
        size = publish(ctx, book, path)
        return {'path': path.relative_to(root).as_posix(), 'sheet': sheet.title, 'rows_written': rows, 'columns_written': columns, 'size_bytes': size}
    finally:
        book.close()


def write(ctx, inputs):
    check(ctx)
    if 'values' not in inputs and 'dataset' not in inputs:
        raise PluginError('EXCEL_INPUT_INVALID', '需要 values 或 dataset')
    root, source = path_for(ctx, inputs['path'], existing=True, writing=True)
    _, target = path_for(ctx, inputs.get('output_path', inputs['path']), writing=True)
    with write_guard(target):
        return _write_locked(ctx, inputs, root, source, target)


def _write_locked(ctx, inputs, root, source, target):
    check(ctx)
    same = source == target
    if same and not inputs.get('overwrite', False):
        raise PluginError('EXCEL_OVERWRITE_REQUIRED', '原地写入需要 overwrite=true，也可通过 output_path 另存')
    if not same and target.exists():
        raise PluginError('EXCEL_TARGET_EXISTS')
    content = file_bytes(source)
    book = load(content, writing=True)
    try:
        sheet = worksheet(book, inputs.get('sheet'), inputs.get('create_sheet', False))
        rows, columns = write_cells(ctx, sheet, matrix(ctx, inputs), inputs.get('start_cell', 'A1'), inputs.get('formulas', False))
        size = publish(ctx, book, target, expected=hashlib.sha256(content).digest() if same else None)
        return {'path': target.relative_to(root).as_posix(), 'sheet': sheet.title, 'rows_written': rows, 'columns_written': columns, 'size_bytes': size}
    finally:
        book.close()
