"""Real XLSX read/create/write and dataset integration, on temporary files."""
import asyncio
import importlib.util
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from taskweave.core.ports import Scope
from taskweave.core.validation import TaskError
from taskweave.infrastructure.storage import Results, uid
from taskweave.plugins.registry import Registry


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('taskweave_excel'), 'Excel plugin must be installed')
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        results = Results(self.root / 'home', Scope(uid(), uid(), uid(), uid()), Registry())
        self.ctx = SimpleNamespace(environment={'excel_root': str(self.root)}, task_parameters={},
                                   cancelled=lambda: False, datasets=results.datasets)

    def call(self, operation, **inputs):
        from taskweave_excel import ExcelPlugin
        action = ExcelPlugin().actions()['excel.' + operation]
        return asyncio.run(action.execute(self.ctx, inputs))

    def test_create_write_read_and_dataset_export(self):
        self.call('create', path='orders.xlsx', sheet='记录', values=[['编号', '金额'], ['001', 12], ['002', 5]])
        sheets = self.call('inspect', path='orders.xlsx')['sheets']
        self.assertEqual(sheets[0]['name'], '记录')
        result = self.call('read', path='orders.xlsx', sheet='记录')
        ref = result['dataset']
        self.assertEqual(self.ctx.datasets.read(ref)['rows'], [{'编号': '001', '金额': 12}, {'编号': '002', '金额': 5}])
        self.call('create', path='result.xlsx', sheet='结果')
        self.call('write', path='result.xlsx', sheet='结果', start_cell='B2', dataset=ref, overwrite=True)
        reread = self.call('read', path='result.xlsx', sheet='结果', range='B2:C4')
        self.assertEqual(self.ctx.datasets.read(reread['dataset'])['rows'], [{'编号': '001', '金额': 12}, {'编号': '002', '金额': 5}])

    def test_write_preserves_other_cells_styles_formulas_and_plain_text(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '模板'
        sheet['A1'] = '标题'
        sheet['A1'].font = Font(bold=True)
        sheet['E5'] = '=1+2'
        workbook.create_sheet('保留')['A1'] = '保留'
        workbook.save(self.root / 'template.xlsx')
        workbook.close()
        before = (self.root / 'template.xlsx').read_bytes()
        self.call('write', path='template.xlsx', output_path='copy.xlsx', sheet='模板', start_cell='B2', values=[['=1+1', False, None]])
        self.assertEqual((self.root / 'template.xlsx').read_bytes(), before)
        book = load_workbook(self.root / 'copy.xlsx')
        self.addCleanup(book.close)
        self.assertEqual(book['模板']['B2'].data_type, 's')
        self.assertEqual(book['模板']['E5'].value, '=1+2')
        self.assertTrue(book['模板']['A1'].font.bold)
        self.assertEqual(book['保留']['A1'].value, '保留')

    def test_formula_read_modes_and_missing_cached_values(self):
        self.call('create', path='formula.xlsx', values=[['formula'], ['=1+2']], formulas=True)
        with self.assertRaises(TaskError) as raised:
            self.call('read', path='formula.xlsx')
        self.assertEqual(raised.exception.code, 'EXCEL_FORMULA_CACHE_MISSING')
        result = self.call('read', path='formula.xlsx', formula_mode='formulas')
        self.assertEqual(self.ctx.datasets.read(result['dataset'])['rows'], [{'formula': '=1+2'}])

    def test_rejects_implicit_overwrite_escape_duplicates_and_write_failure(self):
        self.call('create', path='safe.xlsx', values=[['x'], [1]])
        original = (self.root / 'safe.xlsx').read_bytes()
        for operation, inputs in [('create', {'path': 'safe.xlsx'}),
                                  ('create', {'path': '../escape.xlsx'}),
                                  ('write', {'path': 'safe.xlsx', 'values': [[2]]}),
                                  ('create', {'path': 'bad.xls'})]:
            with self.subTest(operation=operation, inputs=inputs), self.assertRaises(TaskError):
                self.call(operation, **inputs)
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('disk full')):
            with self.assertRaises(TaskError):
                self.call('write', path='safe.xlsx', values=[[2]], overwrite=True)
        self.assertEqual((self.root / 'safe.xlsx').read_bytes(), original)
        self.assertEqual([p.name for p in self.root.iterdir()], ['safe.xlsx'])
        self.call('create', path='duplicate.xlsx', values=[['x', 'X'], [1, 2]])
        with self.assertRaises(TaskError):
            self.call('read', path='duplicate.xlsx')

    def test_dates_empty_table_new_sheet_and_cancel(self):
        from datetime import date
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['日期', '标记'])
        book.active.append([date(2026, 10, 28), True])
        book.save(self.root / 'dates.xlsx')
        book.close()
        result = self.call('read', path='dates.xlsx')
        row = self.ctx.datasets.read(result['dataset'])['rows'][0]
        self.assertTrue(row['日期'].startswith('2026-10-28'))
        self.assertIs(row['标记'], True)
        self.call('write', path='dates.xlsx', sheet='空表', create_sheet=True, values=[['编号']], overwrite=True)
        empty = self.call('read', path='dates.xlsx', sheet='空表')
        self.assertEqual(self.ctx.datasets.describe(empty['dataset'])['row_count'], 0)
        self.ctx.cancelled = lambda: True
        with self.assertRaises(TaskError):
            self.call('create', path='cancelled.xlsx')
        self.assertFalse((self.root / 'cancelled.xlsx').exists())

    def test_complex_workbook_write_is_rejected_without_losing_source(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        book = Workbook()
        sheet = book.active
        sheet.append([1, 2])
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=1, max_col=2, min_row=1))
        sheet.add_chart(chart, 'A5')
        book.save(self.root / 'chart.xlsx')
        book.close()
        original = (self.root / 'chart.xlsx').read_bytes()
        with self.assertRaises(TaskError) as raised:
            self.call('write', path='chart.xlsx', values=[[3]], overwrite=True)
        self.assertEqual(raised.exception.code, 'EXCEL_COMPLEX_WORKBOOK')
        self.assertEqual((self.root / 'chart.xlsx').read_bytes(), original)

    def test_export_uses_all_rows_and_never_silently_rounds_long_integer(self):
        ref = self.ctx.datasets.create('many', [{'name': 'n', 'type': 'integer'}], [[n] for n in range(1001)])
        self.call('create', path='all.xlsx', dataset=ref)
        back = self.call('read', path='all.xlsx')
        self.assertEqual(back['row_count'], 1001)
        with self.assertRaises(TaskError):
            self.call('create', path='rounded.xlsx', values=[[1234567890123456]])
        self.assertFalse((self.root / 'rounded.xlsx').exists())

    def test_cancellation_during_save_and_external_file_change_preserve_source(self):
        from openpyxl.workbook.workbook import Workbook
        self.call('create', path='source.xlsx', values=[['kept']])
        original = (self.root / 'source.xlsx').read_bytes()
        save = Workbook.save
        def cancel_save(book, stream):
            save(book, stream)
            self.ctx.cancelled = lambda: True
        with patch.object(Workbook, 'save', cancel_save), self.assertRaises(TaskError):
            self.call('write', path='source.xlsx', values=[['changed']], overwrite=True)
        self.assertEqual((self.root / 'source.xlsx').read_bytes(), original)
        self.ctx.cancelled = lambda: False
        def concurrent_save(book, stream):
            save(book, stream)
            (self.root / 'source.xlsx').write_bytes(b'other writer')
        with patch.object(Workbook, 'save', concurrent_save), self.assertRaises(TaskError) as raised:
            self.call('write', path='source.xlsx', values=[['changed']], overwrite=True)
        self.assertEqual(raised.exception.code, 'EXCEL_FILE_CHANGED')
        self.assertEqual((self.root / 'source.xlsx').read_bytes(), b'other writer')

    def test_async_action_cancellation_prevents_late_file_publication(self):
        import threading
        from openpyxl.workbook.workbook import Workbook
        from taskweave_excel import ExcelPlugin
        started, release = threading.Event(), threading.Event()
        original = Workbook.save
        def slow_save(book, stream):
            original(book, stream)
            started.set()
            release.wait(5)
        async def exercise():
            task = asyncio.create_task(ExcelPlugin().actions()['excel.create'].execute(self.ctx, {'path': 'late.xlsx'}))
            try:
                self.assertTrue(await asyncio.to_thread(started.wait, 3))
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)
            finally:
                release.set()
        with patch.object(Workbook, 'save', slow_save):
            asyncio.run(exercise())
        self.assertFalse((self.root / 'late.xlsx').exists())

    def test_same_file_writes_from_two_processes_cannot_lose_changes(self):
        import subprocess
        import sys
        import threading
        from concurrent.futures import ThreadPoolExecutor
        from openpyxl.workbook.workbook import Workbook
        self.call('create', path='shared.xlsx', values=[['old A', 'old B']])
        original = Workbook.save
        started, release = threading.Event(), threading.Event()
        def held_save(book, stream):
            original(book, stream)
            started.set()
            release.wait(5)
        child = '''import asyncio, sys
from types import SimpleNamespace
from taskweave_excel import ExcelPlugin
from taskweave.plugins.sdk import PluginError
ctx=SimpleNamespace(environment={'excel_root':sys.argv[1]},task_parameters={},datasets=None,cancelled=lambda:False)
try:
    asyncio.run(ExcelPlugin().actions()['excel.write'].execute(ctx,{'path':'shared.xlsx','start_cell':'B1','values':[['new B']],'overwrite':True}))
except PluginError as e:
    print(e.code)
else:
    print('WRITTEN')
'''
        with ThreadPoolExecutor(1) as pool, patch.object(Workbook, 'save', held_save):
            first = pool.submit(self.call, 'write', path='shared.xlsx', start_cell='A1', values=[['new A']], overwrite=True)
            try:
                self.assertTrue(started.wait(3))
                second = subprocess.run([sys.executable, '-c', child, str(self.root)], capture_output=True, text=True, timeout=4)
            finally:
                release.set()
            first.result()
        self.assertEqual(second.returncode, 0, second.stderr)
        self.assertEqual(second.stdout.strip(), 'EXCEL_FILE_BUSY')
        self.call('write', path='shared.xlsx', start_cell='B1', values=[['new B']], overwrite=True)
        from openpyxl import load_workbook
        book = load_workbook(self.root / 'shared.xlsx')
        self.addCleanup(book.close)
        self.assertEqual((book.active['A1'].value, book.active['B1'].value), ('new A', 'new B'))


class ExcelWorkerTests(unittest.TestCase):
    def test_real_worker_excel_sql_export_pipeline(self):
        from taskweave.application.service import Application
        with tempfile.TemporaryDirectory() as home, tempfile.TemporaryDirectory() as root, Application(home) as app:
            app.configure_plugin('dataset', True)
            app.configure_plugin('excel', True)
            task = app.repo.create_task('Excel SQL')['task_id']
            app.repo.update_task(task, name='Excel SQL', input_schema={'type': 'object', 'properties': {'excel_root': {'type': 'string', 'default': root}}})
            first = app.repo.save_step(task, {'name': '导入', 'capabilities': ['excel.create', 'excel.read'],
                'step_content': 'async def run(ctx, inputs):\n    await ctx.call("excel.create", {"path":"input.xlsx","values":[["name","amount"],["A",12],["A",8],["B",5]]})\n    result=await ctx.call("excel.read", {"path":"input.xlsx"})\n    return ctx.result(data=result)'})
            second = app.repo.save_step(task, {'name': 'SQL', 'capabilities': ['dataset.query'],
                'bindings': {'dataset': {'ref': {'source': 'step', 'step_id': first['step_id'], 'pointer': '/dataset'}}},
                'step_content': 'async def run(ctx, inputs):\n    ref=await ctx.call("dataset.query", {"sources":{"orders":inputs["dataset"]},"sql":"SELECT name,SUM(amount) AS total FROM orders GROUP BY name ORDER BY name"})\n    return ctx.result(data={"dataset":ref})'})
            third = app.repo.save_step(task, {'name': '导出', 'capabilities': ['excel.create', 'excel.write', 'excel.read', 'dataset.read'],
                'bindings': {'dataset': {'ref': {'source': 'step', 'step_id': second['step_id'], 'pointer': '/dataset'}}},
                'step_content': 'async def run(ctx, inputs):\n    await ctx.call("excel.create", {"path":"output.xlsx"})\n    await ctx.call("excel.write", {"path":"output.xlsx","dataset":inputs["dataset"],"overwrite":True})\n    check=await ctx.call("excel.read", {"path":"output.xlsx"})\n    table=await ctx.call("dataset.read", {"dataset":check["dataset"]})\n    return ctx.result(data=table)'})
            for step in (first, second, third):
                app.confirm_step_manual(step['step_id'], step['content_hash'])
            run = app.create_run(task)
            app.coordinator.start(run['run_id'], uid())
            done = app.coordinator.wait(run['run_id'])
            self.assertEqual(done['status'], 'SUCCEEDED', done['attempts'])
            self.assertEqual(app.repo.read_output(run['run_id'], third['step_id'])['rows'], [{'name': 'A', 'total': 20}, {'name': 'B', 'total': 5}])
